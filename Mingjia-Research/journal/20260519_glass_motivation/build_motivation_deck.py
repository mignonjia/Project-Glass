#!/usr/bin/env python3
"""Build a standalone GlassML motivation draft deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
OUT = OUT_DIR / "glass_motivation_draft_v1.pptx"

SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)

BG = RGBColor(255, 255, 255)
INK = RGBColor(28, 32, 40)
MUTED = RGBColor(92, 101, 116)
BLUE = RGBColor(33, 92, 175)
TEAL = RGBColor(0, 128, 128)
ORANGE = RGBColor(196, 95, 34)
GREEN = RGBColor(60, 130, 80)
RED = RGBColor(184, 63, 63)
LIGHT_BLUE = RGBColor(233, 241, 252)
LIGHT_TEAL = RGBColor(230, 246, 246)
LIGHT_ORANGE = RGBColor(252, 240, 232)
LIGHT_GREEN = RGBColor(235, 247, 238)
LIGHT_GRAY = RGBColor(245, 247, 250)


def add_textbox(slide, x, y, w, h, text="", size=18, bold=False,
                color=INK, align=PP_ALIGN.LEFT, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, x, y, w, h, lines, size=15, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.word_wrap = True
    tf.clear()
    for idx, item in enumerate(lines):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        if isinstance(item, tuple):
            text, bold, color, level = item
        else:
            text, bold, color, level = item, False, INK, 0
        p.level = level
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, title, body, fill=LIGHT_GRAY, accent=BLUE):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(215, 221, 230)
    shape.line.width = Pt(1)
    add_textbox(slide, x + 0.12, y + 0.10, w - 0.24, 0.32, title, size=15,
                bold=True, color=accent)
    add_rich_lines(slide, x + 0.12, y + 0.48, w - 0.24, h - 0.58, body,
                   size=12.5, gap=2)
    return shape


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.42, 0.22, 9.1, 0.42, title, size=24, bold=True)
    if subtitle:
        add_textbox(slide, 0.44, 0.66, 9.0, 0.30, subtitle, size=12.5,
                    color=MUTED)
    line = slide.shapes.add_shape(1, Inches(0.43), Inches(1.02),
                                  Inches(9.08), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(220, 225, 232)
    line.line.color.rgb = RGBColor(220, 225, 232)


def add_footer(slide, refs):
    add_textbox(slide, 0.44, 5.18, 9.05, 0.23, refs, size=7.6, color=MUTED)


def add_takeaway(slide, text, color=BLUE):
    shape = slide.shapes.add_shape(1, Inches(0.48), Inches(4.58),
                                   Inches(9.02), Inches(0.44))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(244, 248, 255)
    shape.line.color.rgb = RGBColor(205, 220, 245)
    add_textbox(slide, 0.60, 4.67, 8.78, 0.22, text, size=12.5, bold=True,
                color=color, align=PP_ALIGN.CENTER)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def slide_1(prs):
    s = blank_slide(prs)
    add_title(s, "GlassML Motivation: Memory Capacity Is Becoming The Substrate Problem",
              "Draft section for proposal discussion")
    add_textbox(s, 0.60, 1.23, 8.8, 0.48,
                "Long-horizon ML is shifting from isolated FLOPs to sustained memory residency and rollout movement.",
                size=19, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_card(s, 0.55, 1.95, 2.85, 1.85, "Pretraining",
             [
                 ("Long context / long video", True, INK, 0),
                 "Activation memory grows with sequence horizon",
                 "Workarounds: checkpointing, SP/CP, shorter clips",
             ], fill=LIGHT_BLUE, accent=BLUE)
    add_card(s, 3.58, 1.95, 2.85, 1.85, "Post-training",
             [
                 ("RL / self-distillation", True, INK, 0),
                 "Repeated rollout generation before each update",
                 "Rollout runtime becomes the training bottleneck",
             ], fill=LIGHT_ORANGE, accent=ORANGE)
    add_card(s, 6.61, 1.95, 2.85, 1.85, "Inference",
             [
                 ("Interactive rollout", True, INK, 0),
                 "Low TTFF, per-frame deadlines, low jitter",
                 "Streaming cache and memory movement dominate",
             ], fill=LIGHT_TEAL, accent=TEAL)
    add_takeaway(s, "The common pressure point: keep longer state local, and move rollout data with low latency.")
    add_footer(s, "Refs: StreamDiffusionV2 (arXiv:2511.07399); OpenAI o1 blog; APRIL RL systems; internal GlassML profiling")


def slide_2(prs):
    s = blank_slide(prs)
    add_title(s, "Pretraining Pays A Memory Tax",
              "Capacity pressure turns into recompute, offload, and fine-grained synchronization")
    add_card(s, 0.55, 1.25, 2.95, 2.55, "Checkpointing",
             [
                 ("Saves activation memory", True, BLUE, 0),
                 "Classic checkpointing: 48 GB -> 7 GB in one reported example",
                 "~30% runtime overhead in the same study",
                 "NeMo: full Transformer-layer recomputation adds ~30% per-layer compute",
             ], fill=LIGHT_BLUE, accent=BLUE)
    add_card(s, 3.65, 1.25, 2.95, 2.55, "CPU / Storage Offload",
             [
                 ("Makes capacity look larger", True, ORANGE, 0),
                 "But moves training onto PCIe / host-memory timelines",
                 "Useful for feasibility; risky for sustained utilization",
                 "Video preprocessing and latent storage add more movement",
             ], fill=LIGHT_ORANGE, accent=ORANGE)
    add_card(s, 6.75, 1.25, 2.95, 2.55, "Memory-Forced Parallelism",
             [
                 ("TP / SP / CP make long samples fit", True, TEAL, 0),
                 "But insert layer-wise collectives",
                 "When overlap fails, sync becomes exposed idle time",
                 "Megatron reports 5x activation-memory reduction with SP + selective recompute",
             ], fill=LIGHT_TEAL, accent=TEAL)
    add_takeaway(s, "Extra memory removes workarounds: less recompute, less offload, fewer fine-grained sync points.")
    add_footer(s, "Refs: Chen et al., Training Deep Nets with Sublinear Memory Cost, 2016; Korthikanti et al., MLSys'23; NVIDIA NeMo/Megatron docs")


def slide_3(prs):
    s = blank_slide(prs)
    add_title(s, "Parallelism Is Not Free When It Is Only Solving Capacity",
              "The problem is exposed synchronization, not only communication bytes")
    add_textbox(s, 0.60, 1.20, 4.10, 0.30, "Memory-forced long-sequence training today", size=14, bold=True, color=BLUE)
    add_card(s, 0.60, 1.58, 4.05, 2.48, "Synchronization chain",
             [
                 "TP / SP / CP split a layer or sequence across devices",
                 "Ranks synchronize after attention / MLP boundaries",
                 "A slow rank stalls its group",
                 "The delay can propagate into PP bubbles and DP finish-time skew",
             ], fill=LIGHT_BLUE, accent=BLUE)
    add_textbox(s, 5.25, 1.20, 4.10, 0.30, "What to measure", size=14, bold=True, color=ORANGE)
    add_card(s, 5.25, 1.58, 4.05, 2.48, "Bubble / exposed time",
             [
                 ("exposed comm = NCCL time not overlapped with compute", True, ORANGE, 0),
                 "Profile: compute, NCCL, idle gaps, rank skew",
                 "Plot vs. SP/CP degree and sequence length",
                 "Avoid claiming a universal power law until measured",
             ], fill=LIGHT_ORANGE, accent=ORANGE)
    add_takeaway(s, "Dual-panel memory reduces memory-forced sharding; remaining sync is coarser and easier to overlap.")
    add_footer(s, "Refs: ByteDance OSDI'25 straggler study; Megatron-Core CP docs; PyTorch profiler communication-overlap guidance")


def slide_4(prs):
    s = blank_slide(prs)
    add_title(s, "World-Model Pretraining Is Long-Context Video Training",
              "Seconds of generated experience quickly become tens of thousands of tokens")
    add_textbox(s, 0.58, 1.28, 3.95, 0.35, "Internal world-model numbers", size=14, bold=True, color=BLUE)
    add_card(s, 0.58, 1.68, 3.95, 2.70, "Token and memory growth",
             [
                 ("1 latent frame ~= 1.5k tokens ~= 0.16 s", True, BLUE, 0),
                 ("1 second video ~= 10k tokens", True, BLUE, 0),
                 "3 latent training: 55 GB peak",
                 "9 latent training: 135 GB peak",
                 "15 latent training: OOM",
             ], fill=LIGHT_BLUE, accent=BLUE)
    add_textbox(s, 5.05, 1.28, 3.95, 0.35, "Current coping mechanisms", size=14, bold=True, color=TEAL)
    add_card(s, 5.05, 1.68, 3.95, 2.70, "Make it fit somehow",
             [
                 "Train on short clips, then hope inference generalizes",
                 "Checkpoint long attention / DiT blocks",
                 "Use SP/CP for long rollouts",
                 "Use sliding-window / compressed KV for long video",
                 "Trade stability and horizon for memory feasibility",
             ], fill=LIGHT_TEAL, accent=TEAL)
    add_takeaway(s, "World-model pretraining needs memory capacity before it needs another isolated FLOP headline.")
    add_footer(s, "Refs: internal Mingjia profiling; StreamDiffusionV2; long-video generation / Causal DiT literature")


def slide_5(prs):
    s = blank_slide(prs)
    add_title(s, "Post-Training And Inference Are Both Rollout-Limited",
              "RL turns generation into repeated trajectories; deployment turns it into an interactive loop")
    add_card(s, 0.55, 1.26, 4.30, 2.85, "Post-training: RL / self-distillation",
             [
                 ("Rollout generation dominates modern RL systems", True, ORANGE, 0),
                 "LLM RL systems report rollout generation can exceed 90% of RL runtime",
                 "DeepSeek-R1-style GRPO samples many responses per update",
                 "Video RL is heavier: multi-frame, multi-step denoising rollouts",
                 "World-model reward pipelines add inverse dynamics / preference models",
             ], fill=LIGHT_ORANGE, accent=ORANGE)
    add_card(s, 5.15, 1.26, 4.30, 2.85, "Inference: online planning rollout",
             [
                 ("Agents repeatedly query future trajectories", True, TEAL, 0),
                 "state/action -> predicted future -> policy/planner -> next action",
                 "StreamDiffusionV2: low TTFF + strict per-frame deadlines",
                 "Reported: <0.5 s first frame; 58.28 FPS with 14B model on 4x H100",
                 "Short chunks make memory / communication harder to hide",
             ], fill=LIGHT_TEAL, accent=TEAL)
    add_takeaway(s, "Rollout acceleration is not only compute: it is memory residency, cache movement, and low-jitter sync.")
    add_footer(s, "Refs: APRIL active partial rollouts; DeepSeek-R1; OpenAI o1 blog; StreamDiffusionV2; video/world-model RL post-training papers")


def slide_6(prs):
    s = blank_slide(prs)
    add_title(s, "Why Dual-Panel Glass Belongs In This Trend",
              "A larger local memory domain attacks the root cause before routing details")
    add_card(s, 0.55, 1.25, 2.85, 2.95, "Today: memory-limited node",
             [
                 "HBM capacity wall",
                 "Gradient checkpointing",
                 "CPU/offload",
                 "TP/SP/CP just to fit",
                 "Small rollout batches",
                 "Short cached horizon",
             ], fill=LIGHT_GRAY, accent=RED)
    add_card(s, 3.58, 1.25, 2.85, 2.95, "Dual-panel glass memory",
             [
                 ("GPU + memory panel", True, BLUE, 0),
                 ("Memory-only panel", True, BLUE, 0),
                 "Larger local memory pool",
                 "Longer sequence/sample stays local",
                 "Fewer memory-forced sync points",
                 "Larger rollout batch / horizon",
             ], fill=LIGHT_BLUE, accent=BLUE)
    add_card(s, 6.61, 1.25, 2.85, 2.95, "Quantitative hook",
             [
                 ("192 GB/slot -> ~384 GB/slot", True, GREEN, 0),
                 "Kimi-K2-scale 4 TB model:",
                 "21 memory slots -> 11 memory slots",
                 "~48% fewer memory slots",
                 "Junda proxy: training speedup single 1.04x, double 2.09x, 3-panel 4.17x",
             ], fill=LIGHT_GREEN, accent=GREEN)
    add_takeaway(s, "Glass is not only an interconnect story; it is a memory-capacity substrate for long-horizon ML.")
    add_footer(s, "Refs: Junda dual-sided glass analysis; TCPMT'14/'16 dual-sided glass interposers; internal GlassML panel simulations")


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_1(prs)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs)
    slide_6(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

