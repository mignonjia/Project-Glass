#!/usr/bin/env python3
"""Build v3 motivation deck following flow.md."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Glass ML Presentation Tajana Samsung SK.pptx"
OUT = Path(__file__).resolve().parent / "glass_motivation_draft_v3_flow.pptx"

INK = RGBColor(20, 20, 20)
MUTED = RGBColor(90, 90, 90)
BLUE = RGBColor(31, 78, 121)


def clear_slides(prs):
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        r_id = slide_id.rId
        prs.part.drop_rel(r_id)
        slide_id_list.remove(slide_id)


def set_text_style(paragraph, size=14, bold=False, color=INK):
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_body_lines(placeholder, lines, size=13.4):
    tf = placeholder.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(lines):
        text, level, bold = item
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.space_after = Pt(3)
        set_text_style(p, size=size, bold=bold)


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(0.34), Inches(5.16), Inches(9.10), Inches(0.25))
    tf = box.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = text
    set_text_style(p, size=7.3, color=MUTED)


def add_section(prs, title, subtitle, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.34), Inches(3.05), Inches(9.10), Inches(0.70))
    tf = box.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    p = tf.paragraphs[0]
    p.text = subtitle
    p.alignment = PP_ALIGN.LEFT
    set_text_style(p, size=18, color=BLUE)
    if footer:
        add_footer(slide, footer)
    return slide


def add_title_and_body(prs, title, lines, footer=None, body_size=13.2):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = title
    add_body_lines(slide.placeholders[1], lines, size=body_size)
    if footer:
        add_footer(slide, footer)
    return slide


def add_two_column(prs, title, left_lines, right_lines, footer=None, body_size=12.8):
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    slide.shapes.title.text = title
    add_body_lines(slide.placeholders[1], left_lines, size=body_size)
    add_body_lines(slide.placeholders[2], right_lines, size=body_size)
    if footer:
        add_footer(slide, footer)
    return slide


def main():
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)

    # Slide 1
    add_section(
        prs,
        "Memory Has Become a Big Problem in World Models",
        "From pretraining, to post-training, to inference: every phase increasingly depends on long-horizon rollout state.",
        footer="Draft v3 follows journal/20260519_glass_motivation/flow.md",
    )

    # Slide 2
    add_title_and_body(
        prs,
        "Pretraining: Foundation Video Models Are Natively Long-Context Training",
        [
            ("Video/world-model pretraining is not just bigger images; it is long spatiotemporal context", 0, True),
            ("1 latent frame ≈ 1.5k tokens ≈ 0.16 s; 1 second video ≈ 10k tokens", 0, False),
            ("Internal profile: 3 latent training = 55 GB peak; 9 latent = 135 GB peak; 15 latent = OOM", 0, False),
            ("Reported video-generation training utilization is low: StepVideo Table 2 reports ≈35.9% MFU in a 540P pretraining setting", 0, False),
            ("Memory workaround #1: gradient checkpointing", 0, True),
            ("Reduces peak training memory by not storing many intermediate activations", 1, False),
            ("Classic result reports total memory 48 GB → 7 GB with ~30% runtime overhead", 1, False),
            ("Transformer caveat: total-memory savings depend on how much peak memory is activations; NeMo reports ~30% extra compute for full layer recompute", 1, False),
            ("Memory workaround #2: SP / CP / TP / PP", 0, True),
            ("Makes the sample fit, but adds synchronization points and straggler propagation", 1, False),
            ("More local memory directly targets both problems: fewer checkpointed activations, less memory-forced parallelism", 0, True),
        ],
        footer="Refs: StepVideo technical report (arXiv:2502.10248); Chen et al. 2016; NVIDIA NeMo/Megatron docs; internal Mingjia profiling",
        body_size=12.6,
    )

    # Slide 3
    add_two_column(
        prs,
        "Post-Training: Distillation and RL Both Need Rollout",
        [
            ("Why post-training matters", 0, True),
            ("Step distillation is the efficiency path: convert expensive video diffusion into fewer-step causal generation", 0, False),
            ("Your pipeline: bidirectional finetune (50 steps) → diffusion forcing causal model (50 steps) → self-forcing causal model (4 steps)", 0, False),
            ("RL / reward post-training is the alignment path: action following, preference, robot/world-model reward", 0, False),
            ("LLM trend: reasoning models increasingly scale with RL train-time compute", 0, False),
            ("LLM systems trend: rollout generation can dominate RL runtime (>80-90% reported in recent systems work)", 0, False),
        ],
        [
            ("Why this is a memory problem", 0, True),
            ("Both distillation and RL repeatedly generate trajectories before an update", 0, False),
            ("Video rollout is heavier than text rollout: many frames, multi-step denoising, rolling KV/cache state", 0, False),
            ("Self-forcing loads multiple models; your notes flag memory capacity as a blocker", 0, False),
            ("Longer rollout distillation: 9-latent teacher supervising 18-latent (or longer) student rollouts", 0, False),
            ("Internal measurement hook: Wan1.3B causal, B200, SP={1,4}; measure one denoising step at the last block where KV cache is largest", 0, False),
            ("Glass target: support larger rollout batches / longer horizons with less offload and less sequence sharding", 0, True),
        ],
        footer="Refs: OpenAI o1 blog; DeepSeek-R1 / GRPO; APRIL active partial rollouts; video RL / inverse-dynamics reward literature; internal SP rollout profiling",
        body_size=12.2,
    )

    # Slide 4
    add_title_and_body(
        prs,
        "Inference: Interactive World Models Are Streaming Rollout",
        [
            ("Deployment is not offline batch video generation", 0, True),
            ("World-model agents run a loop: state/action → predicted future → policy/planner → next action", 0, False),
            ("This makes inference a continuous rollout workload with strict latency and low jitter requirements", 0, False),
            ("StreamDiffusionV2 frames the systems target clearly", 0, True),
            ("Live video generation needs low time-to-first-frame and per-frame deadlines", 1, False),
            ("Reported: first frame within 0.5 s; 58.28 FPS with a 14B model on 4×H100", 1, False),
            ("It identifies poor multi-GPU scaling: sequence parallelism has unpredictable communication overhead", 1, False),
            ("It also states short-frame chunks push real-time workloads into a memory-bound regime", 1, False),
            ("Your HYWorldPlay/FastVideo note shows the practical pain: with SP=2, a 2.5 s chunk takes ~2.5 min for the first chunk and ~7 min for the second", 0, False),
            ("Takeaway: inference rollout needs memory locality and low-latency movement, not only peak FLOPs", 0, True),
        ],
        footer="Refs: StreamDiffusionV2, arXiv:2511.07399; internal world-model inference notes",
        body_size=12.8,
    )

    # Slide 5
    add_two_column(
        prs,
        "Why Dual-Panel Glass Belongs in This Trend",
        [
            ("What the trend says", 0, True),
            ("Pretraining: long video hits activation capacity; checkpointing and SP/CP are memory workarounds", 0, False),
            ("Post-training: distillation / RL multiply generation cost through repeated rollouts", 0, False),
            ("Inference: interactive rollout is memory-bound under strict SLOs", 0, False),
            ("Across all phases, the common bottleneck is keeping long-horizon state local", 0, True),
        ],
        [
            ("What dual-panel glass changes", 0, True),
            ("One layer: GPU + memory; second layer: memory-only", 0, False),
            ("Dual-sided glass hook from Junda: 192 GB/slot → 384 GB/slot", 0, False),
            ("Thermal hook: glass κ ≈ 1 W/m·K; at 1000 W GPU, <0.6 W leaks to memory side (<1°C rise)", 0, False),
            ("Larger local memory domain for activations, cache state, and rollout batches", 0, False),
            ("Reduces gradient checkpointing pressure during pretraining", 0, False),
            ("Reduces memory-forced SP/CP and its synchronization exposure", 0, False),
            ("Supports longer rollout horizons for post-training and inference", 0, False),
            ("Core proposal: glass substrate is a memory-capacity platform for world models", 0, True),
        ],
        footer="Refs: dual-sided glass packaging literature (TCPMT'14/'16); Junda thermal/dual-panel analysis; ByteDance OSDI'25 straggler study",
        body_size=12.5,
    )

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
