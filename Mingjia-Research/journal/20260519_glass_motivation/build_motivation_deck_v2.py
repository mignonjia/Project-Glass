#!/usr/bin/env python3
"""Build a standalone motivation deck using the GlassML PPTX theme/style."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Glass ML Presentation Tajana Samsung SK.pptx"
OUT = Path(__file__).resolve().parent / "glass_motivation_draft_v2_style_matched.pptx"

INK = RGBColor(20, 20, 20)
MUTED = RGBColor(95, 95, 95)
BLUE = RGBColor(31, 78, 121)


def clear_slides(prs):
    """Remove template slides while keeping masters/layouts/theme."""
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        r_id = slide_id.rId
        prs.part.drop_rel(r_id)
        slide_id_list.remove(slide_id)


def set_text_style(paragraph, size=14, bold=False, color=INK):
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_body_lines(placeholder, lines, size=14):
    tf = placeholder.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(lines):
        text, level, bold = item
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.space_after = Pt(3)
        set_text_style(p, size=size, bold=bold)


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(0.34), Inches(5.17), Inches(9.1), Inches(0.24))
    tf = box.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = text
    set_text_style(p, size=7.5, color=MUTED)


def add_title_and_body(prs, title, lines, footer=None, body_size=13.2):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = title
    add_body_lines(slide.placeholders[1], lines, size=body_size)
    if footer:
        add_footer(slide, footer)
    return slide


def add_two_column(prs, title, left_lines, right_lines, footer=None, body_size=12.7):
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    slide.shapes.title.text = title
    add_body_lines(slide.placeholders[1], left_lines, size=body_size)
    add_body_lines(slide.placeholders[2], right_lines, size=body_size)
    if footer:
        add_footer(slide, footer)
    return slide


def add_section(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.34), Inches(3.10), Inches(9.1), Inches(0.45))
    tf = box.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    p = tf.paragraphs[0]
    p.text = subtitle
    p.alignment = PP_ALIGN.LEFT
    set_text_style(p, size=18, color=BLUE)
    return slide


def main():
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)

    add_section(
        prs,
        "Motivation: Memory Capacity Is Becoming The Substrate Problem",
        "Long-horizon ML is shifting from isolated FLOPs to sustained memory residency and rollout movement.",
    )

    add_title_and_body(
        prs,
        "Training Pays a Memory Tax",
        [
            ("Capacity pressure turns into recompute, offload, and fine-grained synchronization", 0, True),
            ("Long-context LLM / video / MoE training quickly exceeds local HBM capacity", 0, False),
            ("Gradient checkpointing saves activation memory, but recomputes forward activations during backward", 0, False),
            ("Classic checkpointing example: 48 GB → 7 GB, with ~30% runtime overhead", 1, False),
            ("NeMo reports full Transformer-layer recomputation adds ~30% per-layer compute", 1, False),
            ("Megatron reports 5× activation-memory reduction with sequence parallelism + selective recomputation", 1, False),
            ("CPU/offload and storage movement make capacity look larger, but move training onto PCIe / host-memory timelines", 0, False),
            ("Takeaway: extra local memory removes workarounds, not just bytes", 0, True),
        ],
        footer="Refs: Chen et al., 2016; Korthikanti et al., MLSys'23; NVIDIA NeMo/Megatron recomputation docs",
    )

    add_two_column(
        prs,
        "Parallelism Is Not Free When It Is Only Solving Capacity",
        [
            ("Memory-forced TP / SP / CP", 0, True),
            ("Shard a layer, sequence, or context across devices so the sample fits", 0, False),
            ("Insert layer-wise collectives around attention / MLP boundaries", 0, False),
            ("Raw communication volume may be manageable", 0, False),
            ("The real cost is exposed synchronization when overlap fails", 0, True),
            ("Slow ranks can stall TP/CP groups, create PP bubbles, and delay DP synchronization", 0, False),
        ],
        [
            ("How we should measure it", 0, True),
            ("Profile compute kernels, NCCL collectives, idle gaps, and rank skew", 0, False),
            ("Exposed comm = NCCL time not overlapped with useful compute", 0, False),
            ("Sweep SP/CP degree and sequence length", 0, False),
            ("Avoid claiming a universal power-law bubble distribution before measuring it", 0, False),
            ("Dual-panel memory reduces memory-forced sharding; remaining sync is coarser and easier to overlap", 0, True),
        ],
        footer="Refs: ByteDance OSDI'25 straggler study; Megatron-Core context parallelism docs; PyTorch profiler overlap guidance",
    )

    add_title_and_body(
        prs,
        "World-Model Pretraining Is Long-Context Video Training",
        [
            ("Seconds of generated experience quickly become tens of thousands of tokens", 0, True),
            ("Internal profiling hook", 0, True),
            ("1 latent frame ≈ 1.5k tokens ≈ 0.16 s", 1, False),
            ("1 second video ≈ 10k tokens", 1, False),
            ("3 latent training: 55 GB peak; 9 latent training: 135 GB peak; 15 latent training: OOM", 1, False),
            ("Current coping mechanisms", 0, True),
            ("Train on short clips, then hope inference generalizes to long horizons", 1, False),
            ("Checkpoint long attention / DiT blocks", 1, False),
            ("Use SP/CP, sliding windows, or compressed KV/cache selection", 1, False),
            ("Takeaway: world-model pretraining needs memory capacity before another isolated FLOP headline", 0, True),
        ],
        footer="Refs: internal Mingjia profiling; StreamDiffusionV2; long-video generation / Causal DiT literature",
    )

    add_two_column(
        prs,
        "Post-Training and Inference Are Both Rollout-Limited",
        [
            ("Post-training: RL / self-distillation", 0, True),
            ("RL repeatedly generates rollouts before each update", 0, False),
            ("LLM RL systems report rollout generation can exceed 90% of RL runtime", 0, False),
            ("DeepSeek-R1-style GRPO samples many responses per update", 0, False),
            ("Video/world-model rollouts are heavier: multi-frame, multi-step denoising + cache state", 0, False),
            ("This makes rollout latency and memory movement first-order training costs", 0, True),
        ],
        [
            ("Inference: online planning rollout", 0, True),
            ("Agents repeatedly query future trajectories", 0, False),
            ("state/action → predicted future → policy/planner → next action", 0, False),
            ("StreamDiffusionV2: low TTFF, strict per-frame deadline, low jitter", 0, False),
            ("Reported: <0.5 s first frame; 58.28 FPS with 14B model on 4×H100", 0, False),
            ("Short chunks make memory / communication harder to hide", 0, True),
        ],
        footer="Refs: APRIL active partial rollouts; DeepSeek-R1; OpenAI o1 blog; StreamDiffusionV2 (arXiv:2511.07399)",
    )

    add_title_and_body(
        prs,
        "Why Dual-Panel Glass Belongs in This Trend",
        [
            ("The trend is toward larger local memory domains, not only faster isolated accelerators", 0, True),
            ("Today: HBM capacity wall forces checkpointing, offload, TP/SP/CP, small rollout batches, and short cached horizons", 0, False),
            ("Dual-panel glass: GPU + memory panel, plus memory-only panel", 0, False),
            ("Larger local memory pool keeps longer sequence/sample state resident", 1, False),
            ("Fewer memory-forced fine-grained synchronization points", 1, False),
            ("Larger rollout batches and longer horizons for post-training / inference", 1, False),
            ("Quantitative hook", 0, True),
            ("192 GB/slot → ~384 GB/slot; Kimi-K2-scale 4 TB model: 21 memory slots → 11 memory slots (~48% fewer)", 1, False),
            ("Junda proxy result: training speedup single 1.04×, double 2.09×, 3-panel 4.17×", 1, False),
            ("Takeaway: glass is not only an interconnect story; it is a memory-capacity substrate for long-horizon ML", 0, True),
        ],
        footer="Refs: Junda dual-sided glass analysis; TCPMT'14/'16 dual-sided glass interposers; internal GlassML panel simulations",
    )

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
